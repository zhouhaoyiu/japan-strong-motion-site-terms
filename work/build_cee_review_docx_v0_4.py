"""Build a v0.4 review DOCX for the CEE manuscript direction.

This packaging helper converts the v0.4 submission-prose Markdown candidate
into a Word review document and embeds the five CEE figure drafts for
coauthor/domain-expert review. It reuses the v0.3 document styling and figure
layout helpers so the rendered review files remain comparable across versions.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from build_cee_review_docx_v0_3 import (
    OUTPUTS,
    add_bullet,
    add_code_paragraph,
    add_figures,
    add_inline_runs,
    add_markdown_paragraph,
    add_numbered,
    set_cell_margins,
    set_cell_shading,
    set_headers_and_footers,
    style_document,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
REVIEW_DIR = OUTPUTS / "cee_review_v0_4"
INPUT_MD = OUTPUTS / "cee_article_draft_v0_4_submission_prose.md"
OUTPUT_DOCX = REVIEW_DIR / "manuscript" / "cee_combined_manuscript_v0_4_review.docx"


def add_status_table_v04(doc: Document) -> None:
    table = doc.add_table(rows=6, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.8), Inches(4.7)]
    rows = [
        ("Status", "Internal review draft; not a submitted manuscript."),
        ("Main result", "MF2013 D1400/AVS30 site terms improve SA(1.0 s) and SA(3.0 s) residuals much more than PGA."),
        ("New in v0.4", "Submission-prose candidate with numbered references 1-11 and automated citation/guardrail audit."),
        ("Audit result", "v0.4 automated audit: 0 failures, 2 warnings for known placeholders and unresolved submission gates."),
        ("Integrity guardrail", "No claim of guaranteed acceptance, causal proof, full MF2013 reproduction, full Zhao reproduction, or full J-SHIS hazard-map reproduction."),
        ("Remaining gates", "Expert review, Ref. 2 manual verification, final author statements, public repository/DOI, clean-environment reproduction, and data-license review."),
    ]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
        for i, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "F2F4F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
            cell.width = widths[i]
    doc.add_paragraph()


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Frequency-dependent nonergodic site residuals in Japan strong-motion records")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("1F4D78")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Communications Earth & Environment internal review draft v0.4")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note_run = note.add_run("Automated manuscript-hygiene audit passed with known warnings; not submission-ready.")
    note_run.font.size = Pt(10)
    note_run.font.color.rgb = RGBColor(100, 100, 100)


def add_audit_note(doc: Document) -> None:
    doc.add_heading("Review Note", level=1)
    paragraphs = [
        "This v0.4 review document is generated from the submission-prose Markdown candidate. It is intended for coauthor and domain-expert review before journal submission.",
        "The automated v0.4 audit passed reference sequencing, citation range, CEE policy-reference, overclaim phrase, and conservative-limit checks. It still reports known placeholders for authors, affiliations, acknowledgements, funding, related manuscripts, and repository/DOI fields.",
    ]
    for text in paragraphs:
        p = doc.add_paragraph()
        add_inline_runs(p, text)


def add_reference_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)


def add_markdown_body_v04(doc: Document, markdown_text: str) -> None:
    in_code = False
    in_references = False
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            add_markdown_paragraph(doc, " ".join(paragraph_lines).strip())
            paragraph_lines = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            flush_paragraph()
            in_code = not in_code
            continue
        if in_code:
            if line.strip():
                add_code_paragraph(doc, line)
            continue

        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            doc.add_heading(stripped[2:].strip(), level=1)
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            heading = stripped[3:].strip()
            doc.add_heading(heading, level=1)
            in_references = heading.lower() == "references"
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            add_bullet(doc, stripped[2:].strip())
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            if in_references:
                add_reference_paragraph(doc, stripped)
            else:
                add_numbered(doc, re.sub(r"^\d+\.\s+", "", stripped))
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()


def main() -> None:
    doc = Document()
    style_document(doc)
    set_headers_and_footers(doc, "CEE manuscript review draft v0.4")

    add_cover(doc)
    add_status_table_v04(doc)
    doc.add_page_break()

    add_audit_note(doc)
    markdown_text = INPUT_MD.read_text(encoding="utf-8")
    markdown_text = re.sub(
        r"^# Frequency-dependent nonergodic site residuals in Japan strong-motion records\n+",
        "",
        markdown_text,
    )
    add_markdown_body_v04(doc, markdown_text)
    add_figures(doc)
    set_headers_and_footers(doc, "CEE manuscript review draft v0.4")

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
