"""Build a review DOCX for the CEE manuscript package.

This is a packaging helper, not part of the scientific analysis workflow.
It converts the current Markdown manuscript draft into a Word document and
places the four CEE figure drafts after the figure legends for coauthor review.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = PROJECT_ROOT / "outputs"
PACKAGE = OUTPUTS / "cee_submission_package_v0_2"
INPUT_MD = OUTPUTS / "cee_article_draft_v0_2.md"
OUTPUT_DOCX = PACKAGE / "manuscript" / "cee_combined_manuscript_v0_2_review.docx"


FIGURES = [
    (
        "Figure 1",
        OUTPUTS / "figures" / "cee_fig1_workflow_data_evidence.png",
    ),
    (
        "Figure 2",
        OUTPUTS / "figures" / "cee_fig2_mf2013_site_term_gain_ci.png",
    ),
    (
        "Figure 3",
        OUTPUTS / "figures" / "cee_fig3_residual_site_correlation_shift.png",
    ),
    (
        "Figure 4",
        OUTPUTS / "figures" / "cee_fig4_subgroup_robustness.png",
    ),
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.text = "CEE manuscript review draft v0.2"
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)

        footer_p = section.footer.paragraphs[0]
        footer_p.text = "Internal review draft - not submitted"
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)


def add_status_table(doc: Document) -> None:
    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.8), Inches(4.7)]
    rows = [
        ("Status", "Internal review draft; not a submitted manuscript."),
        ("Main result", "MF2013 D1400/AVS30 site terms improve SA(1.0 s) and SA(3.0 s) residuals much more than PGA."),
        ("Integrity guardrail", "No claim of guaranteed acceptance, causal proof, full MF2013 reproduction, or full J-SHIS hazard-map reproduction."),
        ("Remaining gates", "Expert review, reference-manager check, final author statements, public repository/DOI, and data-license review for derived CSV release."),
    ]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
        for i, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "F2F4F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
            cell.width = widths[i]
    doc.add_paragraph()


def add_code_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(9.5)


def add_markdown_paragraph(doc: Document, paragraph: str) -> None:
    p = doc.add_paragraph()
    add_inline_runs(p, paragraph)


def add_inline_runs(paragraph, text: str) -> None:
    tokens = re.split(r"(`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            run.font.size = Pt(10)
        else:
            paragraph.add_run(token)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)


def add_markdown_body(doc: Document, markdown_text: str) -> None:
    in_code = False
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            add_markdown_paragraph(doc, " ".join(paragraph_lines).strip())
            paragraph_lines = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            flush_paragraph()
            in_code = not in_code
            continue
        if in_code:
            if line.strip():
                add_code_paragraph(doc, line)
            continue

        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            doc.add_heading(stripped[2:].strip(), level=1)
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_heading(stripped[3:].strip(), level=1)
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            add_bullet(doc, stripped[2:].strip())
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            add_numbered(doc, re.sub(r"^\d+\.\s+", "", stripped))
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()


def add_figures(doc: Document) -> None:
    doc.add_page_break()
    doc.add_heading("Embedded Figure Drafts", level=1)
    for idx, (label, figure_path) in enumerate(FIGURES):
        if not figure_path.exists():
            continue
        if idx > 0:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        run.add_picture(str(figure_path), width=Inches(6.0))
        cap = doc.add_paragraph(label)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.runs[0].bold = True


def main() -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Frequency-dependent nonergodic site residuals in Japan strong-motion records")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("1F4D78")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Communications Earth & Environment internal review draft v0.2")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)

    add_status_table(doc)
    doc.add_page_break()

    markdown_text = INPUT_MD.read_text(encoding="utf-8")
    # Avoid duplicating the top H1 because the DOCX has its own cover title.
    markdown_text = re.sub(r"^# Article Draft v0\.2 for Communications Earth & Environment\n+", "", markdown_text)
    add_markdown_body(doc, markdown_text)
    add_figures(doc)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
