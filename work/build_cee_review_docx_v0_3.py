"""Build a v0.3 review DOCX for the CEE manuscript direction.

This is a packaging helper, not part of the scientific analysis workflow. It
converts the v0.3 Markdown manuscript draft into a Word review document and
embeds the five CEE figure drafts for coauthor/domain-expert review.
"""

from __future__ import annotations

import re
import struct
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUTS = PROJECT_ROOT / "outputs"
REVIEW_DIR = OUTPUTS / "cee_review_v0_3"
INPUT_MD = OUTPUTS / "cee_article_draft_v0_3.md"
OUTPUT_DOCX = REVIEW_DIR / "manuscript" / "cee_combined_manuscript_v0_3_review.docx"

FIGURES = [
    ("Figure 1", OUTPUTS / "figures" / "cee_fig1_workflow_data_evidence.png"),
    ("Figure 2", OUTPUTS / "figures" / "cee_fig2_mf2013_site_term_gain_ci.png"),
    ("Figure 3", OUTPUTS / "figures" / "cee_fig3_residual_site_correlation_shift.png"),
    ("Figure 4", OUTPUTS / "figures" / "cee_fig4_subgroup_robustness.png"),
    ("Figure 5", OUTPUTS / "figures" / "cee_fig5_zhao2006_external_gmpe_sensitivity.png"),
]

MAX_FIGURE_WIDTH_IN = 9.25
MAX_FIGURE_HEIGHT_IN = 6.0


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_common_section(section, landscape: bool = False) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def style_document(doc: Document) -> None:
    set_common_section(doc.sections[0], landscape=False)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    set_headers_and_footers(doc, "CEE manuscript review draft v0.3")


def set_headers_and_footers(doc: Document, header_text: str) -> None:
    for section in doc.sections:
        header_p = section.header.paragraphs[0]
        header_p.text = header_text
        header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for run in header_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)

        footer_p = section.footer.paragraphs[0]
        footer_p.text = "Internal review draft - not submitted"
        footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in footer_p.runs:
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(100, 100, 100)


def add_status_table(doc: Document) -> None:
    table = doc.add_table(rows=5, cols=2)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(1.8), Inches(4.7)]
    rows = [
        ("Status", "Internal review draft; not a submitted manuscript."),
        ("Main result", "MF2013 D1400/AVS30 site terms improve SA(1.0 s) and SA(3.0 s) residuals much more than PGA."),
        ("New in v0.3", "Zhao et al. (2006) OpenQuake audit and Figure 5 test whether the long-period D1400 residual pattern is visible outside MF2013."),
        ("Integrity guardrail", "No claim of guaranteed acceptance, causal proof, full MF2013 reproduction, full Zhao reproduction, or full J-SHIS hazard-map reproduction."),
        ("Remaining gates", "Expert review, reference-manager check, final author statements, public repository/DOI, and data-license review for derived CSV release."),
    ]
    for row, (left, right) in zip(table.rows, rows):
        row.cells[0].text = left
        row.cells[1].text = right
        for i, cell in enumerate(row.cells):
            set_cell_margins(cell)
            if i == 0:
                set_cell_shading(cell, "F2F4F7")
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
            cell.width = widths[i]
    doc.add_paragraph()


def add_code_paragraph(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    run.font.name = "Courier New"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
    run.font.size = Pt(9.5)


def add_inline_runs(paragraph, text: str) -> None:
    tokens = re.split(r"(`[^`]+`)", text)
    for token in tokens:
        if not token:
            continue
        if token.startswith("`") and token.endswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Courier New"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Courier New")
            run.font.size = Pt(10)
        else:
            paragraph.add_run(token)


def add_markdown_paragraph(doc: Document, paragraph: str) -> None:
    p = doc.add_paragraph()
    add_inline_runs(p, paragraph)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)


def add_numbered(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(4)
    add_inline_runs(p, text)


def add_markdown_body(doc: Document, markdown_text: str) -> None:
    in_code = False
    paragraph_lines: list[str] = []

    def flush_paragraph() -> None:
        nonlocal paragraph_lines
        if paragraph_lines:
            add_markdown_paragraph(doc, " ".join(paragraph_lines).strip())
            paragraph_lines = []

    for raw_line in markdown_text.splitlines():
        line = raw_line.rstrip()

        if line.startswith("```"):
            flush_paragraph()
            in_code = not in_code
            continue
        if in_code:
            if line.strip():
                add_code_paragraph(doc, line)
            continue

        stripped = line.strip()
        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith("# "):
            flush_paragraph()
            doc.add_heading(stripped[2:].strip(), level=1)
            continue
        if stripped.startswith("## "):
            flush_paragraph()
            doc.add_heading(stripped[3:].strip(), level=1)
            continue
        if stripped.startswith("### "):
            flush_paragraph()
            doc.add_heading(stripped[4:].strip(), level=2)
            continue
        if stripped.startswith("- "):
            flush_paragraph()
            add_bullet(doc, stripped[2:].strip())
            continue
        if re.match(r"^\d+\.\s+", stripped):
            flush_paragraph()
            add_numbered(doc, re.sub(r"^\d+\.\s+", "", stripped))
            continue

        paragraph_lines.append(stripped)

    flush_paragraph()


def png_dimensions(path: Path) -> tuple[int, int] | None:
    """Return PNG pixel dimensions without adding a Pillow dependency."""
    with path.open("rb") as handle:
        header = handle.read(24)
    if len(header) < 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    width, height = struct.unpack(">II", header[16:24])
    return width, height


def figure_width(path: Path) -> Inches:
    dimensions = png_dimensions(path)
    if dimensions is None:
        return Inches(MAX_FIGURE_WIDTH_IN)

    width_px, height_px = dimensions
    if width_px <= 0 or height_px <= 0:
        return Inches(MAX_FIGURE_WIDTH_IN)

    aspect = width_px / height_px
    width_limited_by_height = MAX_FIGURE_HEIGHT_IN * aspect
    return Inches(min(MAX_FIGURE_WIDTH_IN, width_limited_by_height))


def add_figures(doc: Document) -> None:
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    set_common_section(section, landscape=True)
    set_headers_and_footers(doc, "CEE manuscript review draft v0.3")

    doc.add_heading("Embedded Figure Drafts", level=1)
    for idx, (label, figure_path) in enumerate(FIGURES):
        if not figure_path.exists():
            continue
        if idx > 0:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_after = Pt(2)
        run = p.add_run()
        run.add_picture(str(figure_path), width=figure_width(figure_path))
        cap = doc.add_paragraph(label)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_together = True
        cap.paragraph_format.space_before = Pt(0)
        cap_run = cap.runs[0]
        cap_run.bold = True
        cap_run.font.size = Pt(11)


def main() -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run("Frequency-dependent nonergodic site residuals in Japan strong-motion records")
    title_run.bold = True
    title_run.font.size = Pt(18)
    title_run.font.color.rgb = RGBColor.from_string("1F4D78")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Communications Earth & Environment internal review draft v0.3")
    subtitle_run.italic = True
    subtitle_run.font.size = Pt(11)

    add_status_table(doc)
    doc.add_page_break()

    markdown_text = INPUT_MD.read_text(encoding="utf-8")
    markdown_text = re.sub(r"^# Article Draft v0\\.3 for Communications Earth & Environment\n+", "", markdown_text)
    add_markdown_body(doc, markdown_text)
    add_figures(doc)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
